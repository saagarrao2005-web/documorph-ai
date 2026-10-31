import os
import io
import json
import zipfile
from io import BytesIO
from typing import Optional

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, JSONResponse
from fpdf import FPDF
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image
import pypdf

app = FastAPI(title="DocuMorph Python Backend", version="1.0.0")

# Setup CORS — defaults to "*" so every Vercel preview/production domain works
# without extra config. Set DOCUMORPH_ALLOWED_ORIGINS to a comma-separated list
# of your exact frontend origins (e.g. https://documorph-ai.vercel.app) to restrict.
allowed_origins_env = os.getenv("DOCUMORPH_ALLOWED_ORIGINS", "*").strip()
if allowed_origins_env == "*":
    allowed_origins = ["*"]
else:
    allowed_origins = [o.strip() for o in allowed_origins_env.split(",") if o.strip()]

app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition", "X-Original-Size", "X-Reduced-Size"],
)

def clean_text(text: str) -> str:
    """Sanitizes text for standard Latin-1 PDF core fonts without throwing encoding exceptions."""
    if not text:
        return ""
    replacements = {
        "\u201c": '"', "\u201d": '"',
        "\u2018": "'", "\u2019": "'",
        "\u2014": " -- ", "\u2013": " - ",
        "\u2022": "* ", "\u2026": "...",
        "\u00a0": " ", "\t": "    "
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.encode("latin-1", errors="replace").decode("latin-1")

@app.get("/")
def read_root():
    return {
        "status": "DocuMorph Python Conversion Backend is Online",
        "endpoints": ["/convert", "/reduce", "/generate", "/health"],
        "version": "1.0.0"
    }

@app.get("/health")
def health_check():
    return {
        "status": "healthy",
        "supported_conversions": {
            "docx": ["PDF", "TXT"],
            "pptx": ["PDF", "DOCX", "TXT"],
            "pdf": ["DOCX", "TXT"],
            "txt": ["PDF", "DOCX"],
            "jpg": ["PDF", "PNG", "WEBP"],
            "jpeg": ["PDF", "PNG", "WEBP"],
            "png": ["PDF", "JPG", "WEBP"],
            "webp": ["PDF", "PNG", "JPG"]
        },
        "supported_reductions": ["PDF", "DOCX", "PPTX"]
    }

@app.post("/convert")
async def convert_file(
    file: UploadFile = File(...),
    target_format: Optional[str] = Form(None),
    to: Optional[str] = Query(None)
):
    """
    Receives an uploaded file (DOCX, PPTX, PDF, Images, TXT) and converts it
    to the requested target format (PDF, DOCX, TXT, PNG, JPG, WebP).
    """
    try:
        raw_target = target_format or to or "PDF"
        target = raw_target.strip().upper()
        filename = file.filename or "document"
        name_part, ext_part = os.path.splitext(filename)
        src_ext = ext_part.lower().lstrip(".")
        
        contents = await file.read()
        if len(contents) == 0:
            raise HTTPException(status_code=400, detail="The uploaded file is empty.")
        
        # 50 MB limit
        if len(contents) > 50 * 1024 * 1024:
            raise HTTPException(status_code=400, detail="File exceeds maximum size limit of 50 MB.")

        out_buf = BytesIO()
        out_ext = target.lower()
        media_type = "application/octet-stream"

        # 1. Word Document (.docx / .doc)
        if src_ext in ["docx", "doc"]:
            try:
                doc = Document(BytesIO(contents))
            except Exception:
                raise HTTPException(status_code=400, detail="Could not read Word document. Please ensure it is a valid .docx file.")
            
            if target == "PDF":
                pdf = FPDF(orientation="portrait", unit="mm", format="A4")
                pdf.set_auto_page_break(auto=True, margin=15)
                pdf.add_page()
                pdf.set_margins(15, 15, 15)
                
                # Document title header if available
                title_text = clean_text(name_part.replace("-", " ").replace("_", " ").title())
                pdf.set_font("Helvetica", style="B", size=18)
                pdf.cell(0, 12, title_text, new_x="LMARGIN", new_y="NEXT")
                pdf.set_draw_color(180, 180, 190)
                pdf.line(15, pdf.get_y(), 195, pdf.get_y())
                pdf.ln(6)

                pdf.set_font("Helvetica", size=11)
                for paragraph in doc.paragraphs:
                    txt = clean_text(paragraph.text.strip())
                    if not txt:
                        continue
                    if paragraph.style.name.startswith("Heading 1"):
                        pdf.ln(3)
                        pdf.set_font("Helvetica", style="B", size=14)
                        pdf.multi_cell(0, 8, txt)
                        pdf.set_font("Helvetica", size=11)
                        pdf.ln(2)
                    elif paragraph.style.name.startswith("Heading"):
                        pdf.ln(2)
                        pdf.set_font("Helvetica", style="B", size=12)
                        pdf.multi_cell(0, 7, txt)
                        pdf.set_font("Helvetica", size=11)
                        pdf.ln(1)
                    else:
                        pdf.multi_cell(0, 6, txt)
                        pdf.ln(2)
                
                # Render tables
                for table in doc.tables:
                    pdf.ln(2)
                    for row in table.rows:
                        row_text = " | ".join(clean_text(cell.text.strip()) for cell in row.cells if cell.text.strip())
                        if row_text:
                            pdf.set_font("Helvetica", size=10)
                            pdf.multi_cell(0, 5, row_text)
                    pdf.ln(2)

                out_buf.write(pdf.output())
                media_type = "application/pdf"
                out_ext = "pdf"

            elif target == "TXT":
                lines = [p.text for p in doc.paragraphs if p.text.strip()]
                text_content = "\n\n".join(lines)
                out_buf.write(text_content.encode("utf-8"))
                media_type = "text/plain; charset=utf-8"
                out_ext = "txt"
            else:
                raise HTTPException(status_code=400, detail=f"Cannot convert Word (.docx) to {target}. Supported targets: PDF, TXT.")

        # 2. PowerPoint Presentation (.pptx / .ppt)
        elif src_ext in ["pptx", "ppt"]:
            try:
                prs = Presentation(BytesIO(contents))
            except Exception:
                raise HTTPException(status_code=400, detail="Could not read PowerPoint presentation. Please ensure it is a valid .pptx file.")

            if target == "PDF":
                pdf = FPDF(orientation="landscape", unit="mm", format="A4")
                pdf.set_auto_page_break(auto=True, margin=15)
                
                for i, slide in enumerate(prs.slides):
                    pdf.add_page()
                    # Decorative top accent bar
                    pdf.set_fill_color(139, 124, 255)
                    pdf.rect(0, 0, 297, 6, "F")
                    
                    # Slide title
                    title = ""
                    if slide.shapes.title and slide.shapes.title.text:
                        title = clean_text(slide.shapes.title.text.strip())
                    
                    pdf.set_font("Helvetica", style="B", size=20)
                    pdf.set_text_color(30, 20, 40)
                    pdf.set_y(18)
                    pdf.cell(0, 12, title if title else f"Slide {i + 1}", new_x="LMARGIN", new_y="NEXT")
                    pdf.ln(4)
                    
                    # Slide body text & bullet points
                    pdf.set_font("Helvetica", size=13)
                    pdf.set_text_color(60, 55, 75)
                    for shape in slide.shapes:
                        if shape == slide.shapes.title:
                            continue
                        if shape.has_text_frame:
                            for p in shape.text_frame.paragraphs:
                                txt = clean_text(p.text.strip())
                                if txt:
                                    pdf.multi_cell(0, 8, f"{txt}")
                                    pdf.ln(2)
                        elif shape.has_table:
                            for row in shape.table.rows:
                                row_txt = " | ".join(clean_text(c.text.strip()) for c in row.cells if c.text.strip())
                                if row_txt:
                                    pdf.set_font("Helvetica", size=11)
                                    pdf.multi_cell(0, 6, row_txt)
                                    pdf.ln(1)
                    
                    # Slide footer number
                    pdf.set_y(192)
                    pdf.set_font("Helvetica", size=10)
                    pdf.set_text_color(150, 150, 160)
                    pdf.cell(0, 10, f"Slide {i + 1} of {len(prs.slides)}", align="R")

                out_buf.write(pdf.output())
                media_type = "application/pdf"
                out_ext = "pdf"

            elif target == "DOCX":
                out_doc = Document()
                out_doc.add_heading(name_part.replace("-", " ").title(), 0)
                for i, slide in enumerate(prs.slides):
                    title = slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.text else f"Slide {i + 1}"
                    out_doc.add_heading(title, level=1)
                    for shape in slide.shapes:
                        if shape == slide.shapes.title:
                            continue
                        if shape.has_text_frame:
                            for p in shape.text_frame.paragraphs:
                                if p.text.strip():
                                    out_doc.add_paragraph(p.text.strip())
                    if i < len(prs.slides) - 1:
                        out_doc.add_page_break()
                out_doc.save(out_buf)
                media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                out_ext = "docx"

            elif target == "TXT":
                lines = []
                for i, slide in enumerate(prs.slides):
                    title = slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.text else f"Slide {i + 1}"
                    lines.append(f"=== {title} ===")
                    for shape in slide.shapes:
                        if shape == slide.shapes.title:
                            continue
                        if shape.has_text_frame:
                            for p in shape.text_frame.paragraphs:
                                if p.text.strip():
                                    lines.append(p.text.strip())
                    lines.append("")
                out_buf.write("\n".join(lines).encode("utf-8"))
                media_type = "text/plain; charset=utf-8"
                out_ext = "txt"

            else:
                raise HTTPException(status_code=400, detail=f"Cannot convert PowerPoint (.pptx) to {target}. Supported targets: PDF, DOCX, TXT.")

        # 3. PDF Document (.pdf)
        elif src_ext == "pdf":
            try:
                reader = pypdf.PdfReader(BytesIO(contents))
            except Exception:
                raise HTTPException(status_code=400, detail="Could not read PDF. Please ensure it is a valid, unencrypted PDF file.")

            if target in ["DOCX", "WORD"]:
                out_doc = Document()
                out_doc.add_heading(name_part.replace("-", " ").title(), 0)
                for i, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text:
                        for line in text.splitlines():
                            if line.strip():
                                out_doc.add_paragraph(line.strip())
                    if i < len(reader.pages) - 1:
                        out_doc.add_page_break()
                out_doc.save(out_buf)
                media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                out_ext = "docx"

            elif target == "TXT":
                text_pages = []
                for i, page in enumerate(reader.pages):
                    t = page.extract_text() or ""
                    text_pages.append(f"--- Page {i + 1} ---\n{t.strip()}")
                out_buf.write("\n\n".join(text_pages).encode("utf-8"))
                media_type = "text/plain; charset=utf-8"
                out_ext = "txt"

            else:
                raise HTTPException(status_code=400, detail=f"Cannot convert PDF to {target}. Supported targets: DOCX, TXT.")

        # 4. Images (.jpg, .jpeg, .png, .webp)
        elif src_ext in ["jpg", "jpeg", "png", "webp"]:
            try:
                im = Image.open(BytesIO(contents))
            except Exception:
                raise HTTPException(status_code=400, detail="Could not read image file. Please provide a valid JPG, PNG, or WebP image.")

            if target == "PDF":
                if im.mode not in ("RGB", "L"):
                    im = im.convert("RGB")
                im.save(out_buf, format="PDF", resolution=100.0)
                media_type = "application/pdf"
                out_ext = "pdf"

            elif target == "PNG":
                im.save(out_buf, format="PNG")
                media_type = "image/png"
                out_ext = "png"

            elif target in ["JPG", "JPEG"]:
                if im.mode not in ("RGB", "L"):
                    im = im.convert("RGB")
                im.save(out_buf, format="JPEG", quality=92, optimize=True)
                media_type = "image/jpeg"
                out_ext = "jpg"

            elif target == "WEBP":
                im.save(out_buf, format="WEBP", quality=90)
                media_type = "image/webp"
                out_ext = "webp"

            else:
                raise HTTPException(status_code=400, detail=f"Cannot convert image to {target}. Supported targets: PDF, PNG, JPG, WebP.")

        # 5. Plain Text (.txt)
        elif src_ext == "txt":
            text_str = contents.decode("utf-8", errors="replace")
            if target == "PDF":
                pdf = FPDF(orientation="portrait", unit="mm", format="A4")
                pdf.set_auto_page_break(auto=True, margin=15)
                pdf.add_page()
                pdf.set_font("Helvetica", size=11)
                for line in text_str.splitlines():
                    pdf.cell(0, 7, text=clean_text(line), new_x="LMARGIN", new_y="NEXT")
                out_buf.write(pdf.output())
                media_type = "application/pdf"
                out_ext = "pdf"

            elif target in ["DOCX", "WORD"]:
                out_doc = Document()
                for line in text_str.splitlines():
                    if line.strip():
                        out_doc.add_paragraph(line)
                out_doc.save(out_buf)
                media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                out_ext = "docx"

            else:
                raise HTTPException(status_code=400, detail=f"Cannot convert TXT to {target}. Supported targets: PDF, DOCX.")

        else:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported input format: .{src_ext}. Supported formats: DOCX, PPTX, PDF, JPG, PNG, WEBP, TXT."
            )

        out_buf.seek(0)
        output_filename = f"{name_part}.{out_ext}"

        return StreamingResponse(
            out_buf,
            media_type=media_type,
            headers={
                "Content-Disposition": f'attachment; filename="{output_filename}"',
                "Access-Control-Expose-Headers": "Content-Disposition"
            }
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Conversion error: {str(e)}")

@app.post("/reduce")
async def reduce_file_size(
    file: UploadFile = File(...),
    target_pct: int = Form(40)
):
    """
    Shrinks PDF, Word (.docx), and PowerPoint (.pptx) documents to target percentage.
    """
    try:
        filename = file.filename or "document"
        name_part, ext_part = os.path.splitext(filename)
        src_ext = ext_part.lower().lstrip(".")
        
        contents = await file.read()
        orig_size = len(contents)
        if orig_size == 0:
            raise HTTPException(status_code=400, detail="The uploaded file is empty.")
        
        pct = max(10, min(90, int(target_pct)))
        out_buf = BytesIO()
        media_type = "application/octet-stream"

        # 1. PDF Reduction
        if src_ext == "pdf":
            try:
                writer = pypdf.PdfWriter(clone_from=BytesIO(contents))
                quality = max(20, min(80, int(pct * 0.9)))
                
                for page in writer.pages:
                    for img_obj in page.images:
                        try:
                            im = img_obj.image
                            if im.mode not in ("RGB", "L"):
                                im = im.convert("RGB")
                            # If image is large, downsample resolution
                            w, h = im.size
                            if max(w, h) > 1200:
                                scale = 1200 / max(w, h)
                                im = im.resize((max(1, round(w * scale)), max(1, round(h * scale))))
                            img_obj.replace(im, quality=quality)
                        except Exception:
                            pass
                    page.compress_content_streams()

                writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
                writer.write(out_buf)
                media_type = "application/pdf"
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"PDF reduction failed: {str(e)}")

        # 2. DOCX and PPTX Reduction (ZIP based package with media images)
        elif src_ext in ["docx", "pptx"]:
            try:
                in_zip = zipfile.ZipFile(BytesIO(contents), "r")
                out_zip = zipfile.ZipFile(out_buf, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9)
                quality = max(25, min(80, int(pct * 0.85)))

                for item in in_zip.infolist():
                    data = in_zip.read(item.filename)
                    fname_lower = item.filename.lower()
                    if fname_lower.endswith((".png", ".jpg", ".jpeg", ".webp")):
                        try:
                            im = Image.open(BytesIO(data))
                            image_format = (im.format or os.path.splitext(item.filename)[1].lstrip(".")).upper()
                            w, h = im.size
                            # Scale down dimensions proportionally if target pct is small
                            scale = max(0.4, min(1.0, (pct / 100.0) ** 0.5))
                            if scale < 0.95:
                                im = im.resize((max(1, round(w * scale)), max(1, round(h * scale))), Image.Resampling.LANCZOS)
                            comp_buf = BytesIO()
                            if image_format in ("JPG", "JPEG"):
                                if im.mode not in ("RGB", "L"):
                                    im = im.convert("RGB")
                                im.save(comp_buf, format="JPEG", quality=quality, optimize=True)
                            elif image_format == "WEBP":
                                im.save(comp_buf, format="WEBP", quality=quality, method=6)
                            else:
                                # Keep the image encoding aligned with the package filename;
                                # replacing PNG bytes with JPEG bytes can corrupt DOCX/PPTX files.
                                im.save(comp_buf, format="PNG", optimize=True)
                            new_data = comp_buf.getvalue()
                            if len(new_data) < len(data):
                                data = new_data
                        except Exception:
                            pass
                    out_zip.writestr(item, data)

                in_zip.close()
                out_zip.close()
                media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document" if src_ext == "docx" else "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            except Exception as e:
                raise HTTPException(status_code=400, detail=f"Document reduction failed: {str(e)}")

        else:
            raise HTTPException(
                status_code=400,
                detail=f"Document reducer supports PDF, DOCX, and PPTX. Given: .{src_ext}"
            )

        reduced_data = out_buf.getvalue()
        # If reduction didn't save bytes (e.g. file is pure text without images), use original
        if len(reduced_data) >= orig_size:
            final_bytes = contents
            final_size = orig_size
        else:
            final_bytes = reduced_data
            final_size = len(reduced_data)

        output_filename = f"{name_part}_reduced.{src_ext}"
        
        return StreamingResponse(
            BytesIO(final_bytes),
            media_type=media_type,
            headers={
                "Content-Disposition": f'attachment; filename="{output_filename}"',
                "X-Original-Size": str(orig_size),
                "X-Reduced-Size": str(final_size),
                "Access-Control-Expose-Headers": "Content-Disposition, X-Original-Size, X-Reduced-Size"
            }
        )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Reduction error: {str(e)}")

@app.post("/generate")
async def generate_document(
    topic: str = Form(...),
    sections: str = Form(...),
    format: str = Form("PDF")
):
    """
    Generates a structured PDF, PPTX presentation, or DOCX document from prompt sections.
    """
    try:
        fmt_choice = format.strip().upper()
        try:
            sec_list = json.loads(sections)
        except Exception:
            sec_list = [{"n": "Overview", "body": sections}]

        clean_topic = clean_text(topic.strip()) or "DocuMorph Document"
        safe_topic_slug = "".join(c for c in topic.strip()[:30] if c.isalnum() or c in "-_ ").strip().replace(" ", "_") or "document"
        out_buf = BytesIO()

        if fmt_choice == "PDF":
            pdf = FPDF(orientation="portrait", unit="mm", format="A4")
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.add_page()
            
            # Title banner
            pdf.set_fill_color(255, 92, 138)
            pdf.rect(0, 0, 210, 8, "F")
            pdf.set_y(20)
            pdf.set_font("Helvetica", style="B", size=20)
            pdf.multi_cell(0, 10, clean_topic)
            pdf.ln(5)

            for sec in sec_list:
                sec_name = clean_text(sec.get("n", ""))
                sec_body = clean_text(sec.get("body", ""))
                pdf.set_font("Helvetica", style="B", size=13)
                pdf.cell(0, 8, sec_name, new_x="LMARGIN", new_y="NEXT")
                pdf.set_font("Helvetica", size=10)
                pdf.multi_cell(0, 5, sec_body)
                pdf.ln(4)

            out_buf.write(pdf.output())
            media_type = "application/pdf"
            filename = f"documorph_{safe_topic_slug}.pdf"

        elif fmt_choice in ["PPT", "PPTX"]:
            prs = Presentation()
            # Title slide
            title_slide = prs.slides.add_slide(prs.slide_layouts[0])
            title_slide.shapes.title.text = topic.strip()
            title_slide.placeholders[1].text = "Generated with DocuMorph AI"

            # Content slides
            for sec in sec_list:
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = sec.get("n", "")
                slide.placeholders[1].text = sec.get("body", "")

            prs.save(out_buf)
            media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            filename = f"documorph_{safe_topic_slug}.pptx"

        else: # Word / DOCX
            doc = Document()
            doc.add_heading(topic.strip(), 0)
            for sec in sec_list:
                doc.add_heading(sec.get("n", ""), level=1)
                doc.add_paragraph(sec.get("body", ""))
            doc.save(out_buf)
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            filename = f"documorph_{safe_topic_slug}.docx"

        out_buf.seek(0)
        return StreamingResponse(
            out_buf,
            media_type=media_type,
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Access-Control-Expose-Headers": "Content-Disposition"
            }
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Generation failed: {str(e)}")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
